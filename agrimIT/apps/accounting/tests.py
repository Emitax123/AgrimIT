"""Smoke + aislamiento de las páginas de finanzas (Fase 3 del rediseño).

Verifican que balance/cobranzas/account_form renderizan (templates nuevos) y que
las cuentas por cobrar respetan el aislamiento por usuario (regla inviolable §1).
"""
from decimal import Decimal

import pytest
from django.urls import reverse

from apps.accounting.models import Account


def _project_with_account(project_factory, user, estimated, advance):
    project = project_factory(user=user)
    account = Account.objects.create(
        user=user, estimated=Decimal(estimated), advance=Decimal(advance)
    )
    project.account = account
    project.save()
    return project, account


@pytest.mark.django_db
def test_balance_renders(as_owner):
    resp = as_owner.get(reverse('balance'))
    assert resp.status_code == 200


@pytest.mark.django_db
def test_balance_con_filtro_tipo_y_navegacion(as_owner):
    resp = as_owner.get(reverse('balance'), {'year': 2026, 'month': 3, 'tipo': 'mensura'})
    assert resp.status_code == 200
    assert resp.context['tipo'] == 'mensura'


@pytest.mark.django_db
def test_balance_post_compatibilidad(as_owner):
    resp = as_owner.post(reverse('balance'), {'date': '2026-03'})
    assert resp.status_code == 200
    assert resp.context['month_number'] == 3
    assert resp.context['year'] == 2026


@pytest.mark.django_db
def test_cobranzas_renderiza_con_receivable(as_owner, owner, project_factory):
    _project_with_account(project_factory, owner, '500000', '100000')
    resp = as_owner.get(reverse('accounting_display'))
    assert resp.status_code == 200
    rec = resp.context['receivables']
    assert rec['count'] == 1
    assert rec['total'] == Decimal('400000')
    # recién creado -> bucket "al día"
    assert rec['buckets']['al_dia']['count'] == 1


@pytest.mark.django_db
def test_receivables_aislamiento_por_usuario(as_owner, other, project_factory):
    # La cuenta por cobrar pertenece a `other`, no al usuario logueado (`owner`).
    _project_with_account(project_factory, other, '500000', '0')
    resp = as_owner.get(reverse('accounting_display'))
    assert resp.status_code == 200
    assert resp.context['receivables']['count'] == 0


@pytest.mark.django_db
def test_account_form_get(as_owner, owner, project_factory):
    project = project_factory(user=owner)
    resp = as_owner.get(reverse('accform', kwargs={'pk': project.pk}))
    assert resp.status_code == 200


# --- Fase 4: factura no oficial (§7.7) y exportes (§7.4) ---------------------

@pytest.mark.django_db
def test_invoice_dueno_y_saldo(as_owner, owner, project_factory):
    project, account = _project_with_account(project_factory, owner, '850000', '500000')
    resp = as_owner.get(reverse('invoice', kwargs={'pk': project.pk}))
    assert resp.status_code == 200
    # Saldo = estimated − advance
    assert resp.context['saldo'] == Decimal('350000')
    assert resp.context['total_pagos'] == Decimal('500000')
    # Leyenda obligatoria de documento no oficial (§7.7)
    assert 'no válido como factura oficial' in resp.content.decode()


@pytest.mark.django_db
def test_invoice_aislamiento_por_usuario(as_owner, other, project_factory):
    # El proyecto pertenece a `other`; `owner` (logueado) no debe poder verlo.
    project, _ = _project_with_account(project_factory, other, '500000', '0')
    resp = as_owner.get(reverse('invoice', kwargs={'pk': project.pk}))
    assert resp.status_code == 404


@pytest.mark.django_db
def test_export_balance_xlsx(as_owner):
    import io
    from openpyxl import load_workbook

    resp = as_owner.get(reverse('export_balance_xlsx'), {'year': 2026})
    assert resp.status_code == 200
    assert resp['Content-Type'] == (
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    assert 'attachment' in resp['Content-Disposition']
    wb = load_workbook(io.BytesIO(resp.content))  # workbook válido (no corrupto)
    assert wb.sheetnames == ['Balance 2026']


@pytest.mark.django_db
def test_export_movements_xlsx(as_owner, owner, project_factory):
    import io
    from openpyxl import load_workbook

    _project_with_account(project_factory, owner, '500000', '100000')
    resp = as_owner.get(reverse('export_movements_xlsx'))
    assert resp.status_code == 200
    assert resp['Content-Type'] == (
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    wb = load_workbook(io.BytesIO(resp.content))
    assert 'Movimientos' in wb.sheetnames
    assert 'Por cobrar' in wb.sheetnames


@pytest.mark.django_db
def test_export_balance_pdf_renders(as_owner):
    resp = as_owner.get(reverse('export_balance_pdf'), {'year': 2026})
    assert resp.status_code == 200
