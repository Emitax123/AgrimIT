"""Importación masiva de proyectos desde CSV.

Módulo puro (sin HTTP ni acceso a base de datos): recibe el archivo subido y un
diccionario de clientes del usuario, y devuelve las instancias de ``Project``
válidas (sin guardar) más la lista de errores por fila. La vista se encarga de
persistir las válidas en una transacción.

Reusa ``ProjectForm`` para validar cada fila, de modo que hereda los validadores
catastrales (``solo_digitos``), las choices de ``type`` y la opcionalidad de los
campos de titular/mensura definida en el formulario.
"""

import csv
import io

import openpyxl

from .forms import ProjectForm
from .models import Project

# Líneas que empiezan con este prefijo en el CSV se ignoran (guía/comentarios).
COMMENT_PREFIX = '#'

# Nombres de las hojas de la plantilla Excel.
PROJECTS_SHEET = 'Proyectos'
CLIENTS_SHEET = 'Clientes'

# Valores válidos para mostrar en mensajes y en la guía de la plantilla.
TIPO_VALUES = [c[0] for c in Project.TYPE_CHOICES]
TIPO_MENS_VALUES = [c[0] for c in Project.MENS_CHOICES]

# Columna que identifica al cliente por su ID interno. Se valida aparte porque
# ``ProjectForm`` excluye el campo ``client`` (lo asigna la vista).
CLIENT_COLUMN = 'cliente_id'

# Encabezado CSV -> campo del modelo Project. Única fuente de verdad del mapeo.
COLUMN_MAP = {
    'tipo': 'type',
    'tipo_mensura': 'type_mens',
    'partido': 'partido',
    'partida': 'partida',
    'circunscripcion': 'circ',
    'seccion': 'sect',
    'chacra_num': 'chacra_num',
    'chacra_letra': 'chacra_let',
    'quinta_num': 'quinta_num',
    'quinta_letra': 'quinta_let',
    'fraccion_num': 'fraccion_num',
    'fraccion_letra': 'fraccion_let',
    'manzana_num': 'manzana_num',
    'manzana_letra': 'manzana_let',
    'parcela_num': 'parcela_num',
    'parcela_letra': 'parcela_let',
    'subparcela': 'subparcela',
    'calle': 'street',
    'altura': 'street_num',
    'piso': 'floor',
    'depto': 'dept',
    'titular_nombre': 'titular_name',
    'titular_telefono': 'titular_phone',
    'nro_tramite': 'process_num',
}

# Orden de columnas en la plantilla descargable (tipo y cliente_id primero).
TEMPLATE_HEADERS = [
    'tipo',
    CLIENT_COLUMN,
    'tipo_mensura',
    'partido',
    'partida',
    'circunscripcion',
    'seccion',
    'chacra_num',
    'chacra_letra',
    'quinta_num',
    'quinta_letra',
    'fraccion_num',
    'fraccion_letra',
    'manzana_num',
    'manzana_letra',
    'parcela_num',
    'parcela_letra',
    'subparcela',
    'calle',
    'altura',
    'piso',
    'depto',
    'titular_nombre',
    'titular_telefono',
    'nro_tramite',
]

# Encabezados sin los que el archivo no se puede procesar.
REQUIRED_HEADERS = ['tipo', CLIENT_COLUMN]

# Tope de filas de datos por lote, para evitar abusos.
MAX_ROWS = 1000

# Inverso de COLUMN_MAP: campo del modelo -> encabezado CSV, para que los errores
# nombren la columna del archivo (ej. "tipo") en vez de la etiqueta del modelo.
FIELD_TO_HEADER = {field: header for header, field in COLUMN_MAP.items()}


def _form_errors_to_text(form):
    """Aplana los errores de un ProjectForm a un texto legible por fila.

    Usa el nombre de la columna del archivo (ej. ``tipo``) en vez de la etiqueta
    del modelo, y para ``tipo`` agrega la lista de valores admitidos.
    """
    parts = []
    for field, error_list in form.errors.items():
        column = FIELD_TO_HEADER.get(field, field)
        msg = f"{column}: {' '.join(error_list)}"
        if field == 'type':
            msg += f" (valores válidos: {', '.join(TIPO_VALUES)})"
        parts.append(msg)
    return '; '.join(parts)


def _cell_to_str(value):
    """Convierte una celda de Excel a texto. Los enteros (ej. cliente_id) no
    deben quedar como '1.0'."""
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _rows_from_upload(uploaded_file):
    """Devuelve ``(rows, error)`` con las filas del archivo subido (.xlsx o .csv).

    Detecta el formato por la firma de bytes (los .xlsx son ZIP → empiezan con
    ``PK``). Para Excel lee la hoja "Proyectos" (o la activa). ``error`` es un
    string si no se pudo leer, o ``None`` si todo bien.
    """
    raw = uploaded_file.read()

    if raw[:2] == b'PK':  # archivo Excel (.xlsx)
        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception:
            return [], 'No se pudo leer el archivo Excel. Verificá que sea un .xlsx válido.'
        ws = wb[PROJECTS_SHEET] if PROJECTS_SHEET in wb.sheetnames else wb.active
        rows = [[_cell_to_str(c) for c in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
        return rows, None

    # CSV (UTF-8 con BOM opcional, el que escribe Excel).
    try:
        text = raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        return [], 'No se pudo leer el archivo. Guardalo como CSV (UTF-8) o como Excel (.xlsx).'
    return list(csv.reader(io.StringIO(text))), None


def parse_and_validate(uploaded_file, clients_by_id):
    """Parsea y valida una planilla de proyectos (.xlsx o .csv).

    Args:
        uploaded_file: el archivo subido (objeto file-like con bytes).
        clients_by_id: dict ``{int: Client}`` con los clientes del usuario.

    Returns:
        ``(valid, errors)`` donde
        ``valid`` es ``[(nro_fila, instancia_Project_sin_guardar)]`` con ``client``
        ya asignado (falta ``user``, que pone la vista), y
        ``errors`` es ``[(nro_fila, motivo)]``. Un error a nivel de archivo se
        reporta como ``(0, motivo)``.
    """
    valid = []
    errors = []

    rows, read_error = _rows_from_upload(uploaded_file)
    if read_error:
        return [], [(0, read_error)]

    def _is_comment(row):
        return row and row[0].lstrip().startswith(COMMENT_PREFIX)

    # El encabezado es la primera línea no vacía y no comentada.
    header = None
    header_line = 0
    for idx, row in enumerate(rows, start=1):
        if not row or _is_comment(row) or not any(c.strip() for c in row):
            continue
        header = [h.strip() for h in row]
        header_line = idx
        break

    if header is None:
        return [], [(0, 'El archivo está vacío o no tiene encabezados.')]

    missing = [h for h in REQUIRED_HEADERS if h not in header]
    if missing:
        return [], [(0, f"Faltan columnas requeridas: {', '.join(missing)}. "
                        f"Descargá la plantilla para usar el formato correcto.")]

    data_rows = 0
    for idx in range(header_line, len(rows)):
        row = rows[idx]
        line_num = idx + 1  # número de línea real en el archivo (1-based)

        if _is_comment(row):
            continue

        # Mapear encabezado -> valor (zip tolera filas con columnas de más/menos).
        values = {h: (v or '').strip() for h, v in zip(header, row)}

        # Saltear filas completamente vacías.
        if not any(values.get(h) for h in header):
            continue

        data_rows += 1
        if data_rows > MAX_ROWS:
            errors.append((line_num, f"Se superó el tope de {MAX_ROWS} filas; el resto se ignoró."))
            break

        # Resolver el cliente por ID interno.
        client_raw = values.get(CLIENT_COLUMN, '')
        try:
            client = clients_by_id[int(client_raw)]
        except (ValueError, KeyError):
            errors.append((line_num, f"cliente «{client_raw or '(vacío)'}» no encontrado."))
            continue

        # Construir los datos del proyecto y validarlos con ProjectForm.
        form_data = {
            field: values.get(header, '')
            for header, field in COLUMN_MAP.items()
            if header in values
        }
        form = ProjectForm(form_data)
        if form.is_valid():
            instance = form.save(commit=False)
            instance.client = client
            instance.account = None
            instance.closed = False
            valid.append((line_num, instance))
        else:
            errors.append((line_num, _form_errors_to_text(form)))

    if data_rows == 0 and not errors:
        return [], [(0, 'El archivo no tiene filas de datos.')]

    return valid, errors


def build_template_xlsx(clients):
    """Genera la plantilla Excel: hoja "Proyectos" (para cargar) + hoja
    "Clientes" (referencia ID → Nombre del usuario). Devuelve los bytes del .xlsx.
    """
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = PROJECTS_SHEET
    ws.append(TEMPLATE_HEADERS)

    cs = wb.create_sheet(CLIENTS_SHEET)
    cs.append([CLIENT_COLUMN, 'Nombre'])
    for client in clients:
        cs.append([client.id, client.name])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
